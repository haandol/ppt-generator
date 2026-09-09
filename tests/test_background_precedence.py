"""같은 명세로 HTML·PPTX의 배경 정책과 기존 자동 배경을 검증한다."""

from __future__ import annotations

import base64
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from ppt_generator.interfaces import bg_image_utils
from ppt_generator.interfaces.schemas import DesignSpec, PptxSlideSpec
from ppt_generator.tools.pptx.service import ExportService
from ppt_generator.tools.slides.service import SlidesService


@pytest.mark.parametrize("slide_type", ["title", "closing"])
@pytest.mark.parametrize("solid", [None, "#123456"])
@pytest.mark.parametrize("policy", ["gradient", "none"])
def test_background_policy_preserves_legacy_solid_fields(
    tmp_path: Path, slide_type: str, solid: str | None, policy: str
) -> None:
    spec = PptxSlideSpec(slide_type=slide_type, background_color=solid)
    deck = DesignSpec([spec])
    html = (
        SlidesService()
        .generate_from_design_spec(deck, bg_image_policy=policy)
        .slide_htmls[0]
    )
    single_html = SlidesService.render_single_slide_html(
        0, spec, bg_image_policy=policy
    )
    path = (
        ExportService()
        .export_from_design_spec(deck, tmp_path, bg_image_policy=policy)
        .pptx_path
    )
    slide = Presentation(path).slides[0]
    expected_image = policy == "gradient"
    assert ("data:image/png;base64," in html) is expected_image
    assert ("data:image/png;base64," in single_html) is expected_image
    assert (slide._element.find(f".//{qn('a:blip')}") is not None) is expected_image
    if solid and not expected_image:
        assert f"background-color:{solid}" in html
        assert str(slide.background.fill.fore_color.rgb) == solid[1:]


@pytest.mark.parametrize("policy", ["gradient", "none"])
def test_explicit_image_keeps_precedence_over_solid_and_auto_background(
    tmp_path: Path, policy: str
) -> None:
    image_bytes = bg_image_utils.get_bg_image_bytes("light")
    assert image_bytes is not None
    image_src = "data:image/png;base64," + base64.b64encode(image_bytes).decode()
    spec = PptxSlideSpec(
        slide_type="title",
        background_color="#123456",
        background_image_src=image_src,
        background_image_bytes=image_bytes,
    )
    html = SlidesService.render_single_slide_html(0, spec, bg_image_policy=policy)
    path = (
        ExportService()
        .export_from_design_spec(DesignSpec([spec]), tmp_path, bg_image_policy=policy)
        .pptx_path
    )
    slide = Presentation(path).slides[0]
    blip = slide._element.find(f".//{qn('a:blip')}")
    assert blip is not None
    assert slide.part.related_part(blip.get(qn("r:embed"))).blob == image_bytes
    assert image_src in html


@pytest.mark.parametrize("slide_type", ["title", "closing"])
@pytest.mark.parametrize("solid", [None, "#123456"])
def test_generation_fills_missing_special_background_without_replacing_explicit_color(
    slide_type: str, solid: str | None
) -> None:
    from ppt_generator.tools.design.handlers.generation import _enforce_background_color

    spec = PptxSlideSpec(slide_type=slide_type, background_color=solid)
    result = _enforce_background_color(spec, {"background_color": "#F8F7F3"}, "none", 0)
    assert result.background_color == (solid or "#F8F7F3")
