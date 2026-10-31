"""실제 덱의 노트 갱신, 문맥 무결성, Markdown/HTML/PPTX 일치 검증."""

from __future__ import annotations

import asyncio
import json
from dataclasses import asdict, replace
from html.parser import HTMLParser
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from pptx import Presentation

from ppt_generator.interfaces.llm_output_models import SpeakerScriptOutput
from ppt_generator.interfaces.schemas import DesignSpec, ProjectMetadata, SlideOutline
from ppt_generator.server import create_server
from ppt_generator.tools.design.edit_context import encode_signed_context
from ppt_generator.tools.design.handlers.deps import DesignDeps
from ppt_generator.tools.design.handlers.modification import (
    handle_ingest_slide_edit,
    handle_prepare_slide_edit,
)
from ppt_generator.tools.design.service import DesignService
from ppt_generator.tools.pptx.service import ExportService
from ppt_generator.tools.project.service import ProjectService
from ppt_generator.tools.script.controller import register_script_tools
from ppt_generator.tools.script.service import SpeakerScriptService
from ppt_generator.tools.slides.service import SlidesService
from _helpers import make_slide_spec


@pytest.fixture(params=["generated", "imported"])
def deck(
    request, tmp_path: Path, monkeypatch
) -> tuple[ProjectService, SpeakerScriptService, str, Path]:
    monkeypatch.setattr(
        "ppt_generator.tools.project.service.PPT_GENERATOR_HOME", tmp_path
    )
    projects = ProjectService()
    project_id, directory = projects.resolve_project_dir("speaker-test")
    projects.save_metadata(
        directory,
        ProjectMetadata(
            topic="Queue processing",
            num_slides=2,
            source=request.param,
            audience_type="technical",
            presentation_minutes=12,
            purpose="Explain the supplied design",
        ),
    )
    spec = DesignSpec(
        slides=[
            replace(
                make_slide_spec("Queue"),
                speaker_notes="Original first note.\nKeep its qualification.",
            ),
            replace(make_slide_spec("Worker"), speaker_notes="Original last note."),
        ]
    )
    projects.save_design_spec(directory, spec)
    # Notes-only updates must retain even extension data that the renderer ignores.
    path = directory / "design_spec" / "slide_01.json"
    raw = json.loads(path.read_text())
    raw["extension_metadata"] = {"import_marker": ["preserve", 3]}
    path.write_text(json.dumps(raw))
    if request.param == "generated":
        projects.save_outline(
            directory,
            json.dumps(
                {
                    "slides": [
                        asdict(
                            SlideOutline(
                                title="Queue",
                                content_summary="Supplied queue design",
                                speaker_notes=spec.slides[0].speaker_notes,
                                source_refs=["provided-diagram"],
                                assumptions=["Latency has not been measured"],
                            )
                        ),
                        asdict(
                            SlideOutline(
                                title="Worker",
                                content_summary="Worker reads the queue",
                                speaker_notes=spec.slides[1].speaker_notes,
                            )
                        ),
                    ]
                }
            ),
        )
    return projects, SpeakerScriptService(), project_id, directory


def _submission(notes: list[str], *, reverse: bool = False) -> str:
    scripts = [
        {"slide_index": i + 1, "speaker_notes": text} for i, text in enumerate(notes)
    ]
    return json.dumps({"scripts": list(reversed(scripts)) if reverse else scripts})


def _files(directory: Path) -> dict[str, bytes]:
    return {
        p.relative_to(directory).as_posix(): p.read_bytes()
        for p in directory.rglob("*")
        if p.is_file() and p.name != ".edit.lock"
    }


def test_prepare_passes_actual_deck_brief_notes_and_same_response_model(deck) -> None:
    projects, service, project_id, directory = deck
    before = _files(directory)
    prepared = service.prepare(
        projects, project_id, "Korean narration; host takes questions"
    )
    data = json.loads(prepared["user_prompt"])
    assert data["brief"] == "Korean narration; host takes questions"
    assert data["presentation_context"]["presentation_minutes"] == 12
    assert (
        data["slides"][0]["design_spec"]["textboxes"][0]["paragraphs"][0]["runs"][0][
            "text"
        ]
        == "Queue"
    )
    assert data["slides"][0]["design_spec"]["speaker_notes"].startswith(
        "Original first note."
    )
    if data["presentation_context"]["source"] == "generated":
        assert data["slides"][0]["outline"]["source_refs"] == ["provided-diagram"]
        assert data["slides"][0]["outline"]["assumptions"] == [
            "Latency has not been measured"
        ]
    else:
        assert all(slide["outline"] is None for slide in data["slides"])
    assert prepared["response_schema"] == SpeakerScriptOutput.model_json_schema()
    assert prepared["script_context"]
    assert _files(directory) == before


def test_ingest_updates_only_notes_and_preserves_unaffected_notes(deck) -> None:
    projects, service, project_id, directory = deck
    prepared = service.prepare(projects, project_id)
    data = json.loads(prepared["user_prompt"])
    notes = [
        "The queue holds pending work.\nThe worker reads it next.",
        data["slides"][1]["design_spec"]["speaker_notes"],
    ]
    result = service.ingest(
        projects,
        project_id,
        _submission(notes, reverse=True),
        prepared["script_context"],
    )
    for i, slide in enumerate(data["slides"], 1):
        after = json.loads(
            (directory / "design_spec" / f"slide_{i:02d}.json").read_text()
        )
        assert after == {**slide["design_spec"], "speaker_notes": notes[i - 1]}
        if slide["outline"] is not None:
            after_outline = json.loads(projects.load_outline_slide(directory, i - 1))
            assert after_outline == {**slide["outline"], "speaker_notes": notes[i - 1]}
    markdown = Path(result["script_path"]).read_text()
    assert notes[0] in markdown and notes[1] in markdown
    assert markdown.index(notes[0]) < markdown.index(notes[1])
    assert result["export_required"] is True
    if data["presentation_context"]["source"] == "imported":
        assert not (directory / "outline").exists()


class _NotesParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.notes: list[str] = []

    def handle_starttag(self, tag, attrs) -> None:
        value = dict(attrs).get("data-speaker-notes")
        if value is not None:
            self.notes.append(value)


def test_notes_and_line_breaks_match_markdown_html_and_pptx(
    deck, tmp_path: Path
) -> None:
    projects, service, project_id, directory = deck
    prepared = service.prepare(projects, project_id)
    notes = [
        "Here's the queue — it holds work.\n\nAnd the worker reads it.\n<Names & IDs> stay qualified.",
        "이제 설명을 마치겠습니다.\n질문은 진행자가 받습니다.",
    ]
    result = service.ingest(
        projects, project_id, _submission(notes), prepared["script_context"]
    )
    saved = projects.load_design_spec(directory)
    html = SlidesService().generate_from_design_spec(saved, bg_image_policy="none")
    actual_html = []
    for document in html.slide_htmls:
        parser = _NotesParser()
        parser.feed(document)
        actual_html.extend(parser.notes)
    pptx = ExportService().export_from_design_spec(
        saved, output_dir=tmp_path / "exports", bg_image_policy="none"
    )
    presentation = Presentation(pptx.pptx_path)
    actual_pptx = [
        slide.notes_slide.notes_text_frame.text for slide in presentation.slides
    ]
    assert actual_html == actual_pptx == notes
    markdown = Path(result["script_path"]).read_text()
    assert all(note in markdown for note in notes)


@pytest.mark.parametrize(
    "scripts",
    [
        [],
        [{"slide_index": 1, "speaker_notes": "one"}],
        [
            {"slide_index": 1, "speaker_notes": "one"},
            {"slide_index": 1, "speaker_notes": "two"},
        ],
        [
            {"slide_index": 0, "speaker_notes": "one"},
            {"slide_index": 2, "speaker_notes": "two"},
        ],
        [
            {"slide_index": 1, "speaker_notes": "one"},
            {"slide_index": 3, "speaker_notes": "two"},
        ],
        [
            {"slide_index": True, "speaker_notes": "one"},
            {"slide_index": 2, "speaker_notes": "two"},
        ],
        [
            {"slide_index": "1", "speaker_notes": "one"},
            {"slide_index": 2, "speaker_notes": "two"},
        ],
        [
            {"slide_index": 1, "speaker_notes": "one", "textboxes": []},
            {"slide_index": 2, "speaker_notes": "two"},
        ],
    ],
)
def test_invalid_submission_does_not_modify_any_artifact(
    deck, scripts: list[dict]
) -> None:
    projects, service, project_id, directory = deck
    prepared = service.prepare(projects, project_id)
    before = _files(directory)
    with pytest.raises(ValueError):
        service.ingest(
            projects,
            project_id,
            json.dumps({"scripts": scripts}),
            prepared["script_context"],
        )
    assert _files(directory) == before


@pytest.mark.parametrize(
    "change", ["visual", "notes", "order", "metadata", "design_intent", "evidence"]
)
def test_changed_deck_rejects_stale_script_before_writing(deck, change: str) -> None:
    projects, service, project_id, directory = deck
    prepared = service.prepare(projects, project_id)
    path = directory / "design_spec" / "slide_01.json"
    raw = json.loads(path.read_text())
    if change == "visual":
        raw["textboxes"][0]["left_px"] += 1
        path.write_text(json.dumps(raw))
    elif change == "notes":
        raw["speaker_notes"] = "Another saved edit"
        path.write_text(json.dumps(raw))
    elif change == "order":
        projects.move_design_spec_slide(directory, 0, 1)
        if (directory / "outline").exists():
            projects.move_outline_slide(directory, 0, 1)
    elif change == "metadata":
        projects.save_metadata(
            directory,
            replace(projects.load_metadata(directory), presentation_minutes=20),
        )
    elif change == "design_intent":
        projects.design_spec_store.save_design_doc_md(
            directory, "Changed design intent"
        )
    elif change == "evidence":
        if not (directory / "outline").exists():
            projects.save_outline(
                directory,
                json.dumps(
                    {
                        "slides": [
                            asdict(
                                SlideOutline(
                                    title=str(i),
                                    content_summary="Supplied source update",
                                    source_refs=["new source"],
                                )
                            )
                            for i in range(2)
                        ]
                    }
                ),
            )
        else:
            outline = json.loads(projects.load_outline_slide(directory, 0))
            outline["source_refs"] = ["a revised source"]
            projects.save_outline_slide(directory, 0, json.dumps(outline))
    before = _files(directory)
    with pytest.raises(ValueError, match="Stale script_context"):
        service.ingest(
            projects,
            project_id,
            _submission(["new", "new"]),
            prepared["script_context"],
        )
    assert _files(directory) == before


def test_invalid_or_cross_project_context_cannot_write(deck) -> None:
    projects, service, project_id, directory = deck
    before = _files(directory)
    for context in [
        "invalid",
        encode_signed_context(
            {"kind": "speaker_script", "version": 1, "project_id": "other-project"},
            directory,
        ),
    ]:
        with pytest.raises(ValueError):
            service.ingest(projects, project_id, _submission(["new", "new"]), context)
    assert _files(directory) == before


def test_io_failure_rolls_back_notes_outline_and_script(deck, monkeypatch) -> None:
    projects, service, project_id, directory = deck
    (directory / "speaker-script.md").write_text("Existing reading document")
    prepared = service.prepare(projects, project_id)
    before = _files(directory)
    original_write = service._write_json
    calls = 0

    def fail_second_write(path: Path, value: dict) -> None:
        nonlocal calls
        calls += 1
        if calls == 2:
            raise OSError("Simulated save failure")
        original_write(path, value)

    monkeypatch.setattr(service, "_write_json", fail_second_write)
    with pytest.raises(OSError, match="Simulated save failure"):
        service.ingest(
            projects,
            project_id,
            _submission(["new", "new"]),
            prepared["script_context"],
        )
    assert _files(directory) == before


def test_reexport_follows_current_slide_order_and_count(deck) -> None:
    projects, service, project_id, directory = deck
    projects.move_design_spec_slide(directory, 0, 1)
    if (directory / "outline").exists():
        projects.move_outline_slide(directory, 0, 1)
    result = service.export(projects, project_id)
    markdown = Path(result["script_path"]).read_text()
    assert markdown.index("Original last note.") < markdown.index(
        "Original first note."
    )
    projects.delete_design_spec_slide(directory, 1)
    if (directory / "outline").exists():
        projects.delete_outline_slide(directory, 1)
    result = service.export(projects, project_id)
    assert result["slide_count"] == 1
    assert "Original first note." not in Path(result["script_path"]).read_text()


def test_silent_slides_and_other_languages_are_preserved(deck) -> None:
    projects, service, project_id, directory = deck
    prepared = service.prepare(projects, project_id)
    service.ingest(
        projects,
        project_id,
        _submission(["", "이 문장은 한국어로 읽습니다."]),
        prepared["script_context"],
    )
    saved = projects.load_design_spec(directory)
    assert [s.speaker_notes for s in saved.slides] == [
        "",
        "이 문장은 한국어로 읽습니다.",
    ]


@pytest.mark.parametrize(
    "supplied_notes", ["", "Explicitly provided replacement notes."]
)
def test_visual_edit_preserves_saved_notes_or_uses_explicit_notes(
    deck, supplied_notes: str
) -> None:
    projects, _, project_id, directory = deck
    original_notes = projects.load_design_spec_slide(directory, 0).speaker_notes
    deps = DesignDeps(project_service=projects, design_service=DesignService())
    prepared = json.loads(
        handle_prepare_slide_edit(
            deps,
            project_id=project_id,
            action="update",
            slide_index=1,
            title="Updated queue heading",
            content_summary="A new heading for the supplied design",
            component_hint="bullets",
            slide_type="title",
            speaker_notes=supplied_notes,
            color_theme="dark",
        )
    )
    payload = {
        "speaker_notes": "Unrequested model-generated narration",
        "textboxes": [
            {
                "component_id": "title.main",
                "left_px": 64,
                "top_px": 64,
                "width_px": 600,
                "height_px": 60,
                "paragraphs": [{"runs": [{"text": "Updated queue heading"}]}],
            }
        ],
    }
    handle_ingest_slide_edit(
        deps,
        project_id=project_id,
        action="update",
        slide_index=1,
        spec_json=json.dumps(payload),
        edit_context=prepared["edit_context"],
        color_theme="dark",
    )
    assert projects.load_design_spec_slide(directory, 0).speaker_notes == (
        supplied_notes or original_notes
    )
    assert json.loads(projects.load_outline_slide(directory, 0))["speaker_notes"] == (
        supplied_notes or original_notes
    )


def test_missing_project_is_not_created(deck) -> None:
    projects, service, _, directory = deck
    with pytest.raises(FileNotFoundError):
        service.prepare(projects, "missing-project")
    assert not (directory.parent / "missing-project").exists()


def test_unfinished_or_misaligned_deck_is_rejected(deck) -> None:
    projects, service, project_id, directory = deck
    (directory / "design_spec" / "slide_01.json").unlink()
    with pytest.raises(ValueError, match="not contiguous"):
        service.prepare(projects, project_id)


def test_controller_and_server_register_the_paired_tools(deck) -> None:
    projects, service, project_id, _ = deck
    functions = {}
    mcp = MagicMock()

    def tool():
        def register(function):
            functions[function.__name__] = function
            return function

        return register

    mcp.tool = tool
    register_script_tools(mcp, service, projects)
    prepared = json.loads(
        functions["prepare_speaker_script"](project_id, brief="English script")
    )
    result = json.loads(
        functions["ingest_speaker_script"](
            project_id,
            _submission(["Hello.", "Thank you."]),
            prepared["script_context"],
        )
    )
    assert Path(result["script_path"]).is_file()
    assert (
        json.loads(functions["export_speaker_script"](project_id))["slide_count"] == 2
    )
    server_tools = {tool.name for tool in asyncio.run(create_server().list_tools())}
    assert set(functions) <= server_tools
